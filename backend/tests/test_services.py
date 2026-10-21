from pathlib import Path

import pytest
from app.config import get_settings
from app.services import answering, retrieval
from app.services.embeddings import embed
from app.services.evaluation import compare_retrieval, evaluate_cases
from app.services.parsing import ParsedPage, chunk_pages, parse_document


@pytest.fixture
def local_services(tmp_path, monkeypatch):
    monkeypatch.setenv("RAG_CITATION_VERIFIER", "numeric")
    for name in ("REDIS_URL", "QDRANT_URL", "LLM_API_KEY", "LLM_MODEL", "RERANKER_MODEL"):
        monkeypatch.setenv("RAG_" + name, "")
    monkeypatch.setenv("RAG_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("RAG_EMBEDDING_PROVIDER", "hash")
    monkeypatch.setenv("RAG_ENVIRONMENT", "development")
    get_settings.cache_clear()
    retrieval.client.cache_clear()
    answering._cache.clear()
    yield tmp_path
    if retrieval.client.cache_info().currsize:
        retrieval.client().close()
    retrieval.client.cache_clear()
    get_settings.cache_clear()


def sources():
    return [
        {"id": "c1", "document_id": "d1", "document_name": "报销规则.md", "knowledge_base_id": "k1", "page": 2,
            "text": "员工出差住宿报销每晚不超过500元。发票必须在30天内提交财务审核。"},
        {"id": "c2", "document_id": "d2", "document_name": "休假规则.md", "knowledge_base_id": "k1", "page": 1,
            "text": "正式员工每年享有10天带薪年假，需提前三个工作日申请。"},
        {"id": "secret", "document_id": "secret-document", "document_name": "他人秘密.txt", "knowledge_base_id": "private", "page": 1,
            "text": "员工出差住宿报销秘密额度99999元，绝不可泄露给其他账户。"},
    ]


def test_qdrant_scope_fusion_and_cleanup(local_services):
    all_sources = sources()
    retrieval.index_chunks(all_sources)
    result = retrieval.retrieve("出差住宿报销上限是多少", all_sources[:2])
    assert result and result[0]["document_id"] == "d1"
    assert all(c["knowledge_base_id"] == "k1" for c in result)
    assert all(c["chunk_id"] != "secret" for c in result)
    assert result[0]["dense_score"] is not None
    assert result[0]["bm25_score"] > 0
    retrieval.delete_document("d1")
    dense = retrieval.retrieve("出差住宿报销", all_sources[:2], strategy="dense")
    assert all(c["document_id"] != "d1" for c in dense)


def test_single_chunk_bm25_has_positive_match(local_services):
    chunk = sources()[:1]
    retrieval.index_chunks(chunk)
    result = retrieval.retrieve("报销", chunk, strategy="bm25")
    assert result and result[0]["document_id"] == "d1"


def test_definition_with_comparison_keeps_both_independent_sources(local_services):
    chunks = [
        {"id": "lead1", "document_id": "release", "document_name": "发布.md", "knowledge_base_id": "k1",
         "text": "发布前置时间是代码提交至生产环境部署完成的时间。"},
        {"id": "lead2", "document_id": "request", "document_name": "需求.md", "knowledge_base_id": "k1",
         "text": "需求前置时间是需求提出至功能交付完成的时间。"},
    ]
    retrieval.index_chunks(chunks)
    result = answering.answer_question("发布前置时间指的是什么，和需求前置时间有什么不同？", chunks)
    assert {c["document_id"] for c in result["citations"]} == {"release", "request"}
    assert "代码提交" in result["answer"] and "需求提出" in result["answer"]


def test_hash_is_deterministic_and_empty_scope_never_searches(local_services):
    assert embed(["你好 RAG"])[0] == embed(["你好 RAG"])[0]
    assert retrieval.retrieve("secret", []) == []


def test_citations_cache_and_document_revision_invalidate(local_services):
    chunks = sources()[:2]
    retrieval.index_chunks(chunks)
    one = answering.answer_question("住宿报销上限", chunks, mode="agent")
    assert one["provider"] == "extractive" and "[1]" in one["answer"]
    assert one["citations"][0]["page"] == 2
    assert not one["cached"]
    assert answering.answer_question("住宿报销上限", chunks, mode="agent")["cached"]
    chunks[0] = {**chunks[0], "text": "住宿报销上限已调整为600元。"}
    retrieval.index_chunks(chunks)
    revised = answering.answer_question("住宿报销上限", chunks, mode="agent")
    assert not revised["cached"] and "600" in revised["answer"]
    assert "500" not in revised["answer"]


def test_unanswerable_empty_and_ragas_require_real_models(local_services):
    response = answering.answer_question("不存在的知识", [])
    assert not response["citations"] and "没有足够证据" in response["answer"]
    with pytest.raises(ValueError, match="Ragas"):
        evaluate_cases([{"question": "q", "reference": "a"}], [], use_ragas=True)
    result = compare_retrieval([{"question": "q", "reference": "a"}], [])
    assert result["relative_hit_rate_change"] is None
    assert result["hybrid"]["metrics"]["hit_at_5"] is None


def test_labeled_evaluation_and_baseline(local_services):
    chunks = sources()[:2]
    retrieval.index_chunks(chunks)
    cases = [{"question": "出差住宿报销", "reference": "500元", "expected_document_ids": ["d1"]}]
    result = evaluate_cases(cases, chunks)
    assert result["metrics"]["hit_at_5"] == 1
    assert "faithfulness" not in result["metrics"]
    assert result["metrics"]["ragas_enabled"] is False
    comparison = compare_retrieval(cases, chunks)
    assert comparison["dense"]["metrics"]["hit_at_5"] == 1
    assert comparison["relative_hit_rate_change"] == 0


def test_multiformat_and_page_preservation(local_services):
    path = local_services / "notice.txt"
    path.write_bytes("住宿报销每晚500元".encode("gb18030"))
    assert "500" in parse_document(path, "notice.txt")[0].text
    html = local_services / "notice.html"
    html.write_text("<script>ignore me</script><h1>报销规则</h1>", encoding="utf-8")
    assert parse_document(html, html.name)[0].text == "报销规则"
    chunks = chunk_pages([ParsedPage("报销规定" * 500, 7), ParsedPage("年假规定", 9)])
    assert chunks[0]["page"] == 7 and chunks[-1]["page"] == 9
    assert max(len(c["text"]) for c in chunks) <= get_settings().chunk_size
    assert [c["ordinal"] for c in chunks] == list(range(len(chunks)))


def test_docx_tables_and_unsupported_files(local_services):
    from docx import Document
    doc = Document()
    doc.add_paragraph("知识库文档")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "额度", "500元"
    path = local_services / "table.docx"
    doc.save(path)
    assert "额度 | 500元" in parse_document(path, path.name)[0].text
    with pytest.raises(ValueError, match="不支持"):
        parse_document(Path("evil.exe"), "evil.exe")
    with pytest.raises(ValueError, match="没有可检索文本"):
        chunk_pages([ParsedPage("   ")])


def test_short_pdf_text_and_blank_page_do_not_require_ocr(local_services, monkeypatch):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    monkeypatch.setenv("RAG_VISION_MODEL", "")
    monkeypatch.setenv("RAG_OCR_ENABLED", "false")
    get_settings.cache_clear()
    writer = PdfWriter()
    page = writer.add_blank_page(300, 300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 100 Td (short text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_blank_page(300, 300)
    path = local_services / "short.pdf"
    writer.write(path)
    pages = parse_document(path, path.name)
    assert pages[0].text.strip() == "short text" and pages[1].text == ""
    assert chunk_pages(pages)[0]["page"] == 1


def test_agent_reranks_once_after_bounded_rewrites(local_services, monkeypatch):
    from app.services import retrieval as module

    monkeypatch.setenv("RAG_RERANKER_MODEL", "stub-model")
    get_settings.cache_clear()
    calls = []
    class StubReranker:
        def predict(self, pairs):
            calls.append(pairs)
            return list(range(len(pairs)))
    monkeypatch.setattr(module, "reranker", lambda name: StubReranker())
    chunks = sources()[:2]
    module.index_chunks(chunks)
    answer = answering.answer_question("报销 量子物理 未知宇宙 怎么计算", chunks, mode="agent")
    assert len(calls) == 1
    assert len([t for t in answer["trace"] if t["step"] == "rewrite"]) <= 2


def test_ragas_preflight_avoids_generation_cost(local_services, monkeypatch):
    from app.services import evaluation

    monkeypatch.setenv("RAG_LLM_PROVIDER", "ollama")
    monkeypatch.setenv("RAG_LLM_MODEL", "local-test")
    get_settings.cache_clear()
    monkeypatch.setattr(evaluation, "answer_question", lambda *a, **k: pytest.fail("invalid evaluation spent generation"))
    with pytest.raises(ValueError, match="reference"):
        evaluation.evaluate_cases([{"question": "q", "reference": ""}], [], True)
    with pytest.raises(ValueError, match="语义"):
        evaluation.evaluate_cases([{"question": "q", "reference": "a"}], [], True)
    assert answering.llm_ready()  # Ollama does not require a dummy API key.


def test_recorded_evaluation_never_regenerates_and_top5_is_exact(local_services, monkeypatch):
    from app.services import evaluation

    monkeypatch.setattr(evaluation, "answer_question", lambda *a, **k: pytest.fail("history was regenerated"))
    citations = [{"document_id": "other", "text": "证据"}] * 5 + [{"document_id": "expected", "text": "第六名"}]
    result = evaluation.evaluate_cases([{"question": "历史问题", "reference": "历史答案", "recorded_answer": "历史答案",
        "recorded_citations": citations, "expected_document_ids": ["expected"], "message_id": "m1"}], [])
    assert result["metrics"]["hit_at_5"] == 0
    assert result["metrics"]["history_replayed_cases"] == 1
    assert result["results"][0]["answer"] == "历史答案"


def test_context_budget_and_history_aware_cache(local_services):
    chunk = sources()[0]
    history, clipped = answering._budget_context("报销", [{"role": "user", "content": "过长历史" * 2000}],
        [{**chunk, "text": "中文证据" * 20000}])
    size = len("报销".encode()) + sum(len(m["content"].encode()) for m in history) + sum(len(c["text"].encode()) for c in clipped)
    assert size < get_settings().ollama_num_ctx - get_settings().max_output_tokens
    assert clipped and "\ufffd" not in clipped[0]["text"]
    assert answering._key("那它呢", [chunk], 5, "standard", [{"role": "user", "content": "年假"}]) != answering._key(
        "那它呢", [chunk], 5, "standard", [{"role": "user", "content": "报销"}])
